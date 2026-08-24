# -*- coding: utf-8 -*-
"""สร้างไฟล์ Excel master จากฐานข้อมูล — 1 แถวต่อเคส + backup อัตโนมัติ
ไฟล์ถูกเขียนใหม่ทั้งไฟล์ทุกครั้ง (กระจกเงาของฐานข้อมูล — ไม่มีแถวซ้ำ/ตกหล่น)
เคสหลักพันใช้เวลาราว 15 วินาที จึงมี schedule_export() ให้เขียนเบื้องหลังแทนการรอ"""
import json
import os
import shutil
import threading
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

import db

APP_DIR = Path(__file__).parent
SCHEMA = json.loads((APP_DIR / "schema" / "sscc_fields.json").read_text(encoding="utf-8"))
CUSTOM_PATH = APP_DIR / "schema" / "custom_fields.json"


def master_dir():
    """โฟลเดอร์เก็บ Excel — อ่าน config สดทุกครั้ง (ผู้ใช้เปลี่ยนที่เก็บได้จากหน้าเว็บ)"""
    cfg = json.loads((APP_DIR / "config.json").read_text(encoding="utf-8"))
    return Path(cfg.get("master_dir") or (APP_DIR / "output"))


def _custom_fields(used_keys):
    """ฟิลด์ รพ. ที่จะเป็นคอลัมน์ใน Excel: ที่เปิดใช้ + ที่ปิดไว้แต่มีข้อมูลเก่าอยู่
    (ปิดฟิลด์แล้วคอลัมน์เดิมยังอยู่ครบพร้อม label — ไม่ตกไปเป็นคอลัมน์รหัสท้ายตาราง)"""
    if not CUSTOM_PATH.exists():
        return []
    data = json.loads(CUSTOM_PATH.read_text(encoding="utf-8"))
    return [f for f in data.get("fields", [])
            if f.get("enabled", True) or f.get("key") in used_keys]


def _label_for_value(field, v):
    """แปลงรหัสเป็นข้อความอ่านได้ (เก็บทั้งรหัสและข้อความใน Excel)
    ค่าแบบเลือกหลายข้อ (checkbox-group) แปลงทีละรหัสก่อนค่อยรวม"""
    opts = field.get("options", []) or []
    if isinstance(v, list):
        return ", ".join(_label_for_value(field, x) for x in v)
    for o in opts:
        if o["v"] == v:
            return o["t"]
    return v


def export_master():
    out_dir = master_dir()
    out_dir.mkdir(parents=True, exist_ok=True)
    out_path = out_dir / "SSCC_master.xlsx"

    meta_cols = ["case_id", "สถานะ", "เลขที่ผู้ป่วย SSCC", "วันที่บันทึก", "วันที่ส่ง SSCC",
                 "ธงคุณภาพข้อมูล", "รับทราบเหตุผล"]

    cases = db.all_cases_full()
    used_keys = {k for c in cases for k in c["data"]}
    fields = [f for f in SCHEMA["fields"]] + _custom_fields(used_keys)
    known = {f.get("sscc") or f["key"] for f in fields}
    extra_keys = sorted({k for c in cases for k in c["data"] if k not in known and not k.startswith("_")})

    # ชีตอ่านได้อยู่หน้าแรก — เปิดไฟล์มาเจอภาษาไทยเลย ส่วนรหัสดิบ (ไว้เทียบกับเว็บ สป.) อยู่ชีตสอง
    wb = Workbook()
    ws2 = wb.active
    ws2.title = "ข้อมูล (อ่านได้)"
    ws = wb.create_sheet("รหัสดิบ SSCC")

    header = meta_cols + [f.get("sscc") or f["key"] for f in fields] + extra_keys
    header_labels = meta_cols + [f["label"] for f in fields] + [k.replace("legacy|", "") for k in extra_keys]
    for ws_, hdr in ((ws2, header_labels), (ws, header)):
        ws_.append(hdr)
        for c in ws_[1]:
            c.font = Font(bold=True)
            c.fill = PatternFill("solid", fgColor="DDEBF7")
        ws_.freeze_panes = "A2"

    for case in cases:
        d = case["data"]
        dq = d.get("_dq") or {}
        meta = [case["id"], case["status"], case.get("sscc_patient_id") or "",
                case.get("created_at") or "", case.get("submitted_at") or "",
                ", ".join(dq.get("flags") or []), (dq.get("ack") or {}).get("reason") or ""]
        raw_row, readable_row = list(meta), list(meta)
        for f in fields:
            key = f.get("sscc") or f["key"]
            v = d.get(key, "")
            readable_row.append(_label_for_value(f, v) if f.get("options") else
                                (", ".join(v) if isinstance(v, list) else v))
            if isinstance(v, list):
                v = ", ".join(v)
            raw_row.append(v)
        for k in extra_keys:
            raw_row.append(d.get(k, ""))
            readable_row.append(d.get(k, ""))
        ws.append(raw_row)
        ws2.append(readable_row)

    for ws_ in (ws, ws2):
        for i in range(1, min(ws_.max_column, 60) + 1):
            ws_.column_dimensions[get_column_letter(i)].width = 14

    # เขียนลงไฟล์ชั่วคราวก่อนแล้วสลับ — ไฟล์ master ไม่มีวันเสียครึ่งๆ กลางๆ แม้โปรแกรมถูกปิดกลางคัน
    tmp = out_path.with_name(out_path.name + ".tmp")
    wb.save(tmp)
    try:
        os.replace(tmp, out_path)
    except PermissionError:
        tmp.unlink(missing_ok=True)
        raise

    # backup รายวัน
    bak_dir = out_dir / "backups"
    bak_dir.mkdir(exist_ok=True)
    bak = bak_dir / f"SSCC_master_{datetime.now():%Y%m%d}.xlsx"
    shutil.copy2(out_path, bak)
    return str(out_path)


_bg_lock = threading.Lock()
_bg_state = {"running": False, "pending": False, "last_error": None}


def schedule_export():
    """สั่ง export เบื้องหลัง — เรียกซ้ำระหว่างที่กำลังเขียนอยู่จะต่อคิวไว้แค่รอบเดียว"""
    with _bg_lock:
        if _bg_state["running"]:
            _bg_state["pending"] = True
            return
        _bg_state["running"] = True
    threading.Thread(target=_bg_worker, daemon=True).start()


def _bg_worker():
    while True:
        try:
            export_master()
            _bg_state["last_error"] = None
        except PermissionError:
            _bg_state["last_error"] = "ไฟล์ Excel เปิดค้างอยู่ — ปิดไฟล์แล้วกด 'สร้าง Excel ใหม่' หรือบันทึกเคสอีกครั้ง"
        except Exception as e:
            _bg_state["last_error"] = f"{type(e).__name__}: {str(e)[:120]}"
        with _bg_lock:
            if _bg_state["pending"]:
                _bg_state["pending"] = False
                continue
            _bg_state["running"] = False
            return


def status():
    """สถานะไฟล์ master สำหรับโชว์หน้ารวม — เวลาอัปเดตล่าสุดอ่านจาก mtime จริงของไฟล์"""
    p = master_dir() / "SSCC_master.xlsx"
    updated = datetime.fromtimestamp(p.stat().st_mtime).strftime("%d/%m %H:%M") if p.exists() else None
    return {"updated": updated, "error": _bg_state["last_error"], "running": _bg_state["running"]}


if __name__ == "__main__":
    print(export_master())
